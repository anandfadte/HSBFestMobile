import os
import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from backend.database import get_db_connection, init_db, is_postgres, sync_postgres_sequences
SEED_FILE_PATH = BASE_DIR / "data" / "events_seed.xlsx"

DEFAULT_SEED_EVENTS = [
    (1, "Code Sprint 2.0", "Technical", "A 3-hour competitive coding hackathon, solo or in pairs", "2026-11-14", "10:00 AM", "Main Auditorium", 100, 0, "💻"),
    (2, "Robo Race", "Technical", "Build and race a line-following robot against the clock", "2026-11-14", "02:00 PM", "Engineering Lab 2", 40, 50, "🤖"),
    (3, "AI Pitch Fest", "Technical", "Pitch an AI-powered product idea to a panel of judges", "2026-11-15", "11:00 AM", "Seminar Hall A", 60, 0, "🧠"),
    (4, "Web Design Rumble", "Technical", "Build a landing page from a brief in 2 hours, judged live", "2026-11-14", "04:00 PM", "Computer Lab 1", 50, 0, "🎨"),
    (5, "Battle of Bands", "Cultural", "Live band competition, 10-minute sets", "2026-11-15", "06:00 PM", "Open Air Theatre", 300, 0, "🎸"),
    (6, "Dance Fiesta", "Cultural", "Group dance competition, all styles welcome", "2026-11-14", "07:00 PM", "Open Air Theatre", 200, 0, "💃"),
    (7, "Nukkad Natak", "Cultural", "Street play competition on a social theme", "2026-11-15", "03:00 PM", "Amphitheatre", 150, 0, "🎭"),
    (8, "Fashion Fusion", "Cultural", "Themed fashion show, teams of 8-12", "2026-11-15", "08:00 PM", "Main Stage", 250, 100, "👗"),
    (9, "Open Mic Night", "Cultural", "Poetry, stand-up, and acoustic sets, 5 minutes each", "2026-11-14", "09:00 PM", "Cafeteria Lawn", 100, 0, "🎤"),
    (10, "Five-a-Side Football", "Sports", "Knockout football tournament", "2026-11-14", "08:00 AM", "Sports Ground", 80, 0, "⚽"),
    (11, "Chess Championship", "Sports", "Individual rapid chess, Swiss format", "2026-11-15", "09:00 AM", "Indoor Games Hall", 32, 0, "♟️"),
    (12, "Valorant Showdown", "Gaming", "5v5 esports tournament, bring your own peripherals", "2026-11-14", "01:00 PM", "Gaming Arena", 60, 0, "🎮"),
    (13, "UI/UX Design Workshop", "Workshop", "Hands-on Figma workshop, no experience needed", "2026-11-14", "11:00 AM", "Design Lab", 40, 0, "✏️"),
    (14, "Intro to Vibe Coding", "Workshop", "Build a working mini-app using an AI coding assistant", "2026-11-15", "10:00 AM", "Computer Lab 1", 45, 0, "⚡"),
]


def seed_events():
    """Seed the events table from data/events_seed.xlsx (or built-in seed) if not already seeded."""
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if events table already has rows
    cursor.execute("SELECT COUNT(*) AS count FROM events")
    row_count = cursor.fetchone()["count"]
    if row_count > 0:
        conn.close()
        return False, f"Events table already contains {row_count} rows. Seeding skipped."

    rows_to_insert = []

    # First attempt: load from data/events_seed.xlsx
    if SEED_FILE_PATH.exists():
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(str(SEED_FILE_PATH), data_only=True)
            if "Events" in workbook.sheetnames:
                sheet = workbook["Events"]
                sheet_rows = list(sheet.iter_rows(values_only=True))
                if len(sheet_rows) >= 2:
                    for row in sheet_rows[1:]:
                        if not row or all(v is None for v in row):
                            continue
                        event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = row[:10]
                        rows_to_insert.append((
                            int(event_id) if event_id is not None else None,
                            str(title or "").strip(),
                            str(category or "").strip(),
                            str(description or "").strip() if description is not None else None,
                            str(event_date or "").strip(),
                            str(event_time or "").strip(),
                            str(venue or "").strip(),
                            int(capacity) if capacity is not None else 0,
                            int(fee) if fee is not None else 0,
                            str(icon or "").strip() if icon is not None else None,
                        ))
        except Exception as e:
            print(f"Warning: Failed to load from Excel seed file ({e}), falling back to built-in seed list.")

    # Fallback to built-in seed rows if Excel loading was not possible
    if not rows_to_insert:
        rows_to_insert = DEFAULT_SEED_EVENTS

    # Insert rows into events table
    inserted_count = 0
    for row in rows_to_insert:
        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            row,
        )
        inserted_count += 1

    conn.commit()

    # Sync PostgreSQL serial sequence so future inserts start at max(id) + 1
    if is_postgres():
        sync_postgres_sequences(conn)

    conn.close()
    return True, f"Successfully seeded {inserted_count} events."


if __name__ == "__main__":
    status, msg = seed_events()
    print(msg)
    conn = get_db_connection()
    events = conn.execute("SELECT id, title, category, venue, capacity, fee, icon FROM events ORDER BY id").fetchall()
    print(f"\nTotal events in database: {len(events)}")
    for e in events:
        icon_str = e['icon'] or ""
        print(f"[{e['id']}] {e['title']} ({e['category']}) - Venue: {e['venue']}, Capacity: {e['capacity']}, Fee: Rs.{e['fee']}")
    conn.close()
